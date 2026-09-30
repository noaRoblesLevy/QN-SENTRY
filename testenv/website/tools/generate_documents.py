"""Generate the public documents of BadSecurityInc with planted metadata.

The documents are the test data for the Document Metadata module (#7, #8) and, through
the author names, for the Breach module (#12). Every value planted here is listed in
testenv/website/README.md, so the expected findings are known in advance.

Run from the repository root:

    pip install -r testenv/website/tools/requirements.txt
    python testenv/website/tools/generate_documents.py

The files are written to testenv/website/files/. All people and data are fictitious.
"""

import io
import re
import shutil
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "files"
COMPANY = "BadSecurityInc"

# ---------- What is planted in which document ----------

PDFS = [
    {
        "file": "company-brochure-2026.pdf",
        "heading": "BadSecurityInc: office IT that just works",
        "text": "Workplace IT, networks, cloud mail and a helpdesk that answers within the hour.",
        "meta": {
            "/Author": "Lotte Van den Broeck",
            "/Title": "BadSecurityInc brochure 2026",
            "/Creator": "Microsoft® Word 2016",
            "/Producer": "Microsoft® Word 2016",
        },
        "created": "2026-01-15",
    },
    {
        "file": "terms-and-conditions.pdf",
        "heading": "General terms and conditions",
        "text": "These terms apply to every quote, order and service agreement of BadSecurityInc.",
        "meta": {
            "/Author": "jpeeters",  # a username instead of a full name
            "/Title": "Microsoft Word - terms and conditions v4.docx",
            "/Creator": "Microsoft® Word 2010",
            "/Producer": "Microsoft® Word 2010",
        },
        "created": "2023-03-02",
    },
    {
        "file": "employee-handbook.pdf",
        "heading": "Employee handbook",
        "text": "Welcome to BadSecurityInc. This handbook explains working hours, leave and IT rules.",
        "meta": {
            "/Author": "Emma Claes",
            # Internal file server path, left in the title by the PDF converter
            "/Title": r"\\SRV-FS01\HR\Handbook\employee-handbook-v3.docx",
            "/Creator": "Microsoft® Word 2010",
            "/Producer": "Microsoft® Word 2010",
        },
        "created": "2024-09-10",
    },
    {
        "file": "it-security-policy.pdf",
        "heading": "IT security policy",
        "text": "Passwords must be at least 8 characters. Remote access is provided through the VPN.",
        "meta": {
            "/Author": "Lars Janssens",
            "/Subject": "Draft, internal use only",
            # A local user profile path reveals the Windows username
            "/Keywords": r"C:\Users\ljanssens\Documents\Policies",
            "/Creator": "Microsoft® Word 2013",
            "/Producer": "Microsoft® Word 2013",
        },
        "created": "2025-05-20",
    },
    {
        "file": "annual-report-2025.pdf",
        "heading": "Annual report 2025",
        "text": "Revenue grew by 12 percent. We welcomed 31 new customers in the Antwerp region.",
        "meta": {
            # Part-time accountant: not on the team page, so only found through metadata
            "/Author": "Pieter Mertens",
            "/Creator": "Microsoft® Excel® 2013",
            "/Producer": "Microsoft® Excel® 2013",
        },
        "created": "2026-03-28",
    },
    {
        "file": "signed-order-form-example.pdf",
        "heading": "Order form (scanned example)",
        "text": "Customer signature: ____________________   Date: ____/____/________",
        "meta": {
            # Scanned on the office printer: model and device name leak
            "/Creator": "BSI-MFP-2F",
            "/Producer": "KONICA MINOLTA bizhub C308",
        },
        "created": "2025-11-04",
    },
    {
        "file": "privacy-notice.pdf",
        "heading": "Privacy notice",
        "text": "We process personal data only to deliver our services. Contact info@badsecurityinc.be.",
        "meta": None,  # Cleaned before publishing: no metadata at all (the good example)
        "created": "2025-02-01",
    },
]

OFFICE = [
    {
        "file": "price-list-2026.xlsx",
        "core": {"creator": "Tom Wouters", "lastModifiedBy": r"BSI\twouters", "title": "Price list 2026"},
        "app": {"Application": "Microsoft Excel", "AppVersion": "15.0300", "Company": COMPANY},
        "rows": [("Service", "Price per month (EUR)"), ("Workplace IT", 45), ("Helpdesk", 25), ("Cloud mail", 6)],
        "created": "2026-01-05",
    },
    {
        "file": "budget-2026.xlsx",
        "core": {"creator": "Pieter Mertens", "lastModifiedBy": "Sofie Maes", "title": "Budget 2026"},
        "app": {
            "Application": "Microsoft Excel",
            "AppVersion": "12.0000",  # Excel 2007: long out of support
            "Company": COMPANY,
            # Relative links in the workbook resolve against the finance share
            "HyperlinkBase": "\\\\SRV-FS01\\Finance\\Budget\\",
        },
        "rows": [("Cost centre", "Budget 2026 (EUR)"), ("Salaries", 412000), ("Hardware", 95000), ("Marketing", 18000)],
        "created": "2025-12-12",
    },
    {
        "file": "delivery-schedule-q4.xlsx",
        "core": {"creator": "Gérard Dubois", "lastModifiedBy": r"BSI\gdubois", "title": "Delivery schedule Q4"},
        "app": {"Application": "Microsoft Excel", "AppVersion": "16.0300", "Company": COMPANY},
        "rows": [("Week", "Customer", "Items"), (41, "Bakkerij (fictitious)", "4 laptops"), (43, "Garage (fictitious)", "1 printer")],
        "created": "2025-09-22",
    },
    {
        "file": "job-vacancy-it-administrator.docx",
        "core": {"author": "Lars Janssens", "last_modified_by": r"BSI\ljanssens", "title": "Vacancy IT administrator"},
        "app": {
            "Application": "Microsoft Office Word",
            "AppVersion": "14.0000",  # Word 2010
            "Company": COMPANY,
            "Template": "\\\\SRV-FS01\\Templates\\BSI-letterhead.dotx",
        },
        "text": [
            "Vacancy: IT administrator (full time)",
            "You manage our Windows servers, the file server and the firewall of our customers.",
            "Apply at jobs@badsecurityinc.be.",
        ],
        "created": "2026-02-18",
    },
    {
        "file": "company-presentation.pptx",
        "core": {"author": "Lotte Van den Broeck", "last_modified_by": "Gérard Dubois", "title": "Company presentation"},
        "app": {"Application": "Microsoft Office PowerPoint", "AppVersion": "16.0000", "Company": COMPANY},
        "slides": ["BadSecurityInc", "Office IT for Belgian SMEs since 2014"],
        "created": "2025-10-01",
    },
]

# ---------- Writers ----------


def utc(day: str) -> datetime:
    return datetime.fromisoformat(day).replace(hour=9, tzinfo=timezone.utc)


def pdf_date(day: str) -> str:
    return utc(day).strftime("D:%Y%m%d%H%M%S+00'00'")


def write_pdf(spec: dict) -> None:
    body = io.BytesIO()
    page = canvas.Canvas(body, pagesize=A4)
    page.setFont("Helvetica-Bold", 20)
    page.drawString(72, 760, spec["heading"])
    page.setFont("Helvetica", 11)
    page.drawString(72, 730, spec["text"])
    page.drawString(72, 90, "BadSecurityInc is a fictitious company used to test QN-Sentry.")
    page.save()

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(body.getvalue())))
    if spec["meta"] is None:
        writer.metadata = None  # remove the document information dictionary completely
    else:
        date = pdf_date(spec["created"])
        writer.metadata = {**spec["meta"], "/CreationDate": date, "/ModDate": date}
    with open(OUT / spec["file"], "wb") as f:
        writer.write(f)


def write_office(spec: dict) -> None:
    path = OUT / spec["file"]
    created = utc(spec["created"])

    if path.suffix == ".xlsx":
        book = Workbook()
        sheet = book.active
        for row in spec["rows"]:
            sheet.append(row)
        book.properties.creator = spec["core"]["creator"]
        book.properties.lastModifiedBy = spec["core"]["lastModifiedBy"]
        book.properties.title = spec["core"]["title"]
        book.properties.created = created.replace(tzinfo=None)
        book.properties.modified = created.replace(tzinfo=None)
        book.save(path)
    elif path.suffix == ".docx":
        doc = Document()
        for line in spec["text"]:
            doc.add_paragraph(line)
        doc.add_paragraph("BadSecurityInc is a fictitious company used to test QN-Sentry.")
        props = doc.core_properties
        props.author = spec["core"]["author"]
        props.last_modified_by = spec["core"]["last_modified_by"]
        props.title = spec["core"]["title"]
        props.created = props.modified = created
        doc.save(path)
    else:
        deck = Presentation()
        for title in spec["slides"]:
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = title
            note = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.5))
            note.text_frame.text = "BadSecurityInc is a fictitious company used to test QN-Sentry."
            note.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        props = deck.core_properties
        props.author = spec["core"]["author"]
        props.last_modified_by = spec["core"]["last_modified_by"]
        props.title = spec["core"]["title"]
        props.created = props.modified = created
        deck.save(path)

    set_extended_properties(path, spec["app"])


APP_XML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
APP_XML_TYPE = "application/vnd.openxmlformats-officedocument.extended-properties+xml"
APP_XML_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties"


def set_extended_properties(path: Path, values: dict[str, str]) -> None:
    """Replace docProps/app.xml: the program, version, company, template and link base.

    The Python libraries cannot set these, but Office writes them into every file.
    """
    fields = "".join(f"<{key}>{escape(value)}</{key}>" for key, value in values.items())
    app_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<Properties xmlns="{APP_XML_NS}" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        f"{fields}</Properties>"
    )

    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}

    entries["docProps/app.xml"] = app_xml.encode()
    content_types = entries["[Content_Types].xml"].decode()
    if "/docProps/app.xml" not in content_types:
        content_types = content_types.replace(
            "</Types>", f'<Override PartName="/docProps/app.xml" ContentType="{APP_XML_TYPE}"/></Types>'
        )
        entries["[Content_Types].xml"] = content_types.encode()
    rels = entries["_rels/.rels"].decode()
    if "docProps/app.xml" not in rels:
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="rIdApp" Type="{APP_XML_REL}" Target="docProps/app.xml"/></Relationships>',
        )
        entries["_rels/.rels"] = rels.encode()

    with tempfile.NamedTemporaryFile(delete=False, suffix=path.suffix) as tmp:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as target:
            # [Content_Types].xml must be the first entry of an Office file
            for name in ["[Content_Types].xml", *[n for n in entries if n != "[Content_Types].xml"]]:
                target.writestr(name, entries[name])
    shutil.move(tmp.name, path)


def escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------- Check what was written ----------


def summary() -> None:
    for path in sorted(OUT.iterdir()):
        if path.suffix == ".pdf":
            info = PdfReader(path).metadata or {}
            values = {k: info[k] for k in info}
        else:
            with zipfile.ZipFile(path) as z:
                core = z.read("docProps/core.xml").decode()
                app = z.read("docProps/app.xml").decode()
            # openpyxl puts an xmlns attribute on each tag, python-docx does not
            values = dict(re.findall(r"<(?:dc:|cp:)?(creator|lastModifiedBy|title)(?:\s[^>]*)?>([^<]*)<", core))
            values.update(re.findall(r"<(Application|AppVersion|Template|HyperlinkBase|Company)>([^<]*)<", app))
        print(f"{path.name}: {values or 'no metadata'}")


def main() -> None:
    OUT.mkdir(exist_ok=True)
    for spec in PDFS:
        write_pdf(spec)
    for spec in OFFICE:
        write_office(spec)
    summary()


if __name__ == "__main__":
    main()
